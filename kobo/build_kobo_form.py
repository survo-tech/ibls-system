#!/usr/bin/env python3
"""
Add the hidden `rid` question to the IBLS XLSForm.

  python3 kobo/build_kobo_form.py SOURCE.xlsx kobo/Self_india_business_leaders_survey_v2_rid.xlsx

* inserts one row  name=rid, type=hidden  at the top level, right after the metadata/calculate rows and
  before the `section_consent` group;
* sets settings.version to 2610200100;
* leaves every other cell, sheet, translation and setting untouched (form_id stays the same so the
  existing Kobo asset / Enketo link https://ee.kobotoolbox.org/jbh8Urj6 is kept when you replace the form).

The survey link then prefills it with ?d[rid]=<RID> (the Worker's /s/<RID> redirect does this).
"""
import sys

from openpyxl import load_workbook

NEW_VERSION = "2610200100"


def main(src: str, dst: str) -> None:
    wb = load_workbook(src)
    ws = wb["survey"]
    header = [c.value for c in ws[1]]
    col = {name: i + 1 for i, name in enumerate(header)}
    names = [ws.cell(r, col["name"]).value for r in range(1, ws.max_row + 1)]
    if "rid" in names:
        raise SystemExit("form already has a 'rid' question")
    target = names.index("section_consent") + 1  # 1-based row of the consent group
    # sanity: everything above it is metadata / calculate rows
    for r in range(2, target):
        t = ws.cell(r, col["type"]).value
        assert t in ("start", "end", "today", "deviceid", "calculate", "username", "phonenumber", "audit"), (r, t)
    ws.insert_rows(target)
    ws.cell(target, col["name"]).value = "rid"
    ws.cell(target, col["type"]).value = "hidden"
    s = wb["settings"]
    shead = [c.value for c in s[1]]
    s.cell(2, shead.index("version") + 1).value = NEW_VERSION
    wb.save(dst)
    print(f"inserted hidden 'rid' at survey row {target}; version {NEW_VERSION}; saved {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
