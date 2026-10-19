#!/usr/bin/env python3
"""
Import contacts into the D1 database `ibls`.

Sources (contact data is personal data: keep it OUT of git, e.g. in ./data/ which is git-ignored):
  --tracker IBLS_outreach_tracker.xlsx   the full tracker (sheet "Master", all 8,092 rows incl. on-hold)
  --csv contacts.csv                     any CSV with at least: email, company; optional rid, salutation,
                                         director_name, cin, state, region, size_band, nature, email_tier,
                                         peer_tier, status (queued|on_hold), notes

What it does
  * computes `send_order`: a stratified-random interleave over region x size_band x nature x email_tier so
    every send day gets a balanced slice of the sample (within each stratum the i-th of n contacts gets key
    (i + U(0,1)) / n; all contacts are then sorted by key);
  * pre-assigns mailboxes round-robin over the five sending mailboxes along that order (invites actually come
    from one shared queue: the Worker sets the real mailbox when a contact is reserved);
  * marks on-hold rows (tracker "On hold - review" / shared consultant) as status on_hold;
  * writes SQL chunks (UPSERT: re-importing never resets progress such as status, invited_at, clicks);
  * optionally runs `npx wrangler d1 execute ibls --local|--remote --file <chunk>` for each chunk.

Precomputed batches/dates in the source are ignored: scheduling is done live by the Worker.

Examples
  python3 scripts/import_contacts.py --tracker data/IBLS_outreach_tracker.xlsx --out build/import --apply --remote
  python3 scripts/import_contacts.py --csv data/new_contacts.csv --order-offset 100000 --apply --remote
"""
from __future__ import annotations

import argparse
import collections
import csv
import os
import random
import string
import subprocess
import sys
from pathlib import Path

MAILBOXES = ["participate@survo.tech", "research@survo.tech", "survey@survo.tech", "study@survo.tech", "team@survo.tech"]
SIZE_TEXT = {"<10cr": "under ₹10 crore", "10-100cr": "₹10–100 crore", ">100cr": "over ₹100 crore"}
RID_ALPHABET = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
COLUMNS = ["rid", "email", "salutation", "director_name", "company", "cin", "state", "region", "size_band", "nature",
           "peer_tier", "email_tier", "mailbox", "status", "send_order", "notes", "is_test"]


def sql_str(v) -> str:
    """SQL literal for a value (NULL for None/empty)."""
    if v is None:
        return "NULL"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return str(int(v))
    s = str(v).strip()
    if s == "":
        return "NULL"
    return "'" + s.replace("'", "''") + "'"


def peer_tier(size_band: str, nature: str) -> str:
    if not size_band or not nature:
        return "similar"
    return f"{SIZE_TEXT.get(size_band, size_band)} {nature.lower()}"


def read_tracker(path: str) -> list[dict]:
    from openpyxl import load_workbook  # only needed for xlsx input

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Master"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}

    def g(r, name):
        i = idx.get(name)
        v = r[i] if i is not None and i < len(r) else None
        return "" if v is None else str(v).strip()

    out = []
    for r in rows:
        if not r or not g(r, "RID"):
            continue
        hold = g(r, "Invite status").lower().startswith("on hold") or g(r, "Shared consultant").upper() == "TRUE"
        out.append({
            "rid": g(r, "RID").upper(), "email": g(r, "Email").lower(), "salutation": g(r, "Salutation") or "Dear Director,",
            "director_name": g(r, "Director name"), "company": g(r, "Company"), "cin": g(r, "CIN"), "state": g(r, "State"),
            "region": g(r, "Region"), "size_band": g(r, "Size band"), "nature": g(r, "Nature"),
            "peer_tier": g(r, "Peer tier (for mail)") or peer_tier(g(r, "Size band"), g(r, "Nature")),
            "email_tier": g(r, "Email tier"), "status": "on_hold" if hold else "queued", "notes": g(r, "Notes"),
        })
    return out


def read_csv(path: str) -> list[dict]:
    out = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f):
            r = {k.strip().lower(): (v or "").strip() for k, v in r.items() if k}
            if not r.get("email") or not r.get("company"):
                continue
            st = r.get("status", "queued").lower()
            out.append({
                "rid": r.get("rid", "").upper(), "email": r["email"].lower(), "salutation": r.get("salutation") or "Dear Director,",
                "director_name": r.get("director_name", ""), "company": r["company"], "cin": r.get("cin", ""), "state": r.get("state", ""),
                "region": r.get("region", ""), "size_band": r.get("size_band", ""), "nature": r.get("nature", ""),
                "peer_tier": r.get("peer_tier") or peer_tier(r.get("size_band", ""), r.get("nature", "")),
                "email_tier": r.get("email_tier", ""), "status": "on_hold" if st in ("on_hold", "hold", "on hold - review") else "queued",
                "notes": r.get("notes", ""),
            })
    return out


def ensure_rids(rows: list[dict], rng: random.Random) -> None:
    used = {r["rid"] for r in rows if r["rid"]}
    for r in rows:
        if not r["rid"]:
            while True:
                rid = "IB" + "".join(rng.choice(RID_ALPHABET) for _ in range(6))
                if rid not in used:
                    used.add(rid)
                    r["rid"] = rid
                    break


def assign_order(rows: list[dict], rng: random.Random, offset: int = 0) -> None:
    """Stratified interleave for sendable rows, then on-hold rows after them; round-robin mailboxes."""
    def stratified(subset):
        strata = collections.defaultdict(list)
        for r in subset:
            strata[(r["region"], r["size_band"], r["nature"], r["email_tier"])].append(r)
        keyed = []
        for key in sorted(strata):
            lst = strata[key]
            rng.shuffle(lst)
            n = len(lst)
            for i, r in enumerate(lst):
                keyed.append(((i + rng.random()) / n, r))
        keyed.sort(key=lambda t: t[0])
        return [r for _, r in keyed]

    ordered = stratified([r for r in rows if r["status"] == "queued"]) + stratified([r for r in rows if r["status"] != "queued"])
    for pos, r in enumerate(ordered):
        r["send_order"] = offset + pos + 1
        r["mailbox"] = MAILBOXES[pos % len(MAILBOXES)]
        r.setdefault("is_test", 0)


def validate(rows: list[dict]) -> list[str]:
    problems = []
    seen_rid, seen_email = set(), set()
    for r in rows:
        if r["rid"] in seen_rid:
            problems.append(f"duplicate rid {r['rid']}")
        if r["email"] in seen_email:
            problems.append(f"duplicate email {r['email']}")
        if "@" not in r["email"]:
            problems.append(f"bad email {r['email']!r} ({r['rid']})")
        if not (4 <= len(r["rid"]) <= 16) or not all(c in string.ascii_uppercase + string.digits for c in r["rid"]):
            problems.append(f"bad rid {r['rid']!r}")
        seen_rid.add(r["rid"])
        seen_email.add(r["email"])
    return problems


def upsert_sql(rows: list[dict], per_statement: int = 50) -> list[str]:
    """Multi-row UPSERT statements. Progress columns are never overwritten on re-import."""
    stmts = []
    for i in range(0, len(rows), per_statement):
        part = rows[i:i + per_statement]
        values = ",\n".join("(" + ",".join(sql_str(r.get(c)) for c in COLUMNS) + ")" for r in part)
        stmts.append(
            f"INSERT INTO contacts ({','.join(COLUMNS)}) VALUES\n{values}\n"
            "ON CONFLICT(rid) DO UPDATE SET salutation=excluded.salutation, director_name=excluded.director_name,"
            " company=excluded.company, cin=excluded.cin, state=excluded.state, region=excluded.region,"
            " size_band=excluded.size_band, nature=excluded.nature, peer_tier=excluded.peer_tier, email_tier=excluded.email_tier,"
            " send_order=CASE WHEN contacts.invited_at IS NULL THEN excluded.send_order ELSE contacts.send_order END,"
            " mailbox=CASE WHEN contacts.invited_at IS NULL THEN excluded.mailbox ELSE contacts.mailbox END;"
        )
    return stmts


def write_chunks(stmts: list[str], out_dir: Path, per_file: int = 10) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("contacts_*.sql"):
        old.unlink()
    files = []
    for n, i in enumerate(range(0, len(stmts), per_file), start=1):
        p = out_dir / f"contacts_{n:03d}.sql"
        p.write_text("\n".join(stmts[i:i + per_file]) + "\n", encoding="utf-8")
        files.append(p)
    return files


def apply_files(files: list[Path], db: str, remote: bool) -> None:
    for p in files:
        cmd = ["npx", "wrangler", "d1", "execute", db, "--remote" if remote else "--local", "--file", str(p), "-y"]
        print("+", " ".join(cmd), flush=True)
        subprocess.run(cmd, check=True)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--tracker", help="IBLS_outreach_tracker.xlsx (sheet Master)")
    src.add_argument("--csv", help="CSV of contacts")
    ap.add_argument("--out", default="build/import", help="directory for SQL chunks")
    ap.add_argument("--db", default="ibls")
    ap.add_argument("--seed", type=int, default=20261027, help="random seed for the stratified order")
    ap.add_argument("--order-offset", type=int, default=0, help="add to send_order (e.g. 100000 to queue a new list after the existing one)")
    ap.add_argument("--apply", action="store_true", help="run wrangler d1 execute for each chunk")
    where = ap.add_mutually_exclusive_group()
    where.add_argument("--remote", action="store_true")
    where.add_argument("--local", action="store_true")
    a = ap.parse_args(argv)

    rng = random.Random(a.seed)
    rows = read_tracker(a.tracker) if a.tracker else read_csv(a.csv)
    ensure_rids(rows, rng)
    problems = validate(rows)
    if problems:
        print("Refusing to import:\n  " + "\n  ".join(problems[:50]), file=sys.stderr)
        return 2
    assign_order(rows, rng, a.order_offset)
    files = write_chunks(upsert_sql(rows), Path(a.out))
    c = collections.Counter(r["status"] for r in rows)
    mb = collections.Counter(r["mailbox"] for r in rows if r["status"] == "queued")
    print(f"{len(rows)} contacts ({dict(c)}); queued per mailbox {dict(mb)}; {len(files)} SQL files in {a.out}")
    if a.apply:
        if not (a.remote or a.local):
            print("--apply needs --remote or --local", file=sys.stderr)
            return 2
        apply_files(files, a.db, a.remote)
    return 0


if __name__ == "__main__":
    sys.exit(main())
